"""
Excel 读取 — 把工作表转换成与 CSV 相同的「表格行」，复用 CSV 导入的全部能力
（表头行自动检测、多行表头合并、末尾合计/签字行识别、类型推断、单事务导入）。

支持企业系统里常见的几种「Excel」：
  - 标准 .xlsx / .xlsm（openpyxl）
  - 样式表不规范的 .xlsx（部分 ERP / 收银系统导出，openpyxl 会报 Fill 错误）→ 直接解析工作表 XML
  - 后缀是 .xls 实际是 XML 的「XML 电子表格 2003」（支付宝账务明细等）
  - 后缀是 .xls 实际是制表符分隔文本（库龄表等报表导出），自动识别 UTF-8 / GBK
  - 真正的旧版二进制 .xls（xlrd）

直接读取单元格的原始值而不是经过 pandas：
  - 日期单元格输出 ISO 格式（2026-05-01 / 2026-05-01 10:30:00），能被识别为 DATE / TIMESTAMP
  - 整数值的浮点数（Excel 里的 3 读出来是 3.0）还原为 3，避免整数列被误判为小数
  - 合并单元格只有左上角有值，其余为空，与 Excel 中看到的一致
  - 公式单元格读取缓存的计算结果，不执行公式
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from typing import Any

EXCEL_TYPES = {"xlsx", "xlsm", "xls"}
MAX_SHEETS = 50
MAX_CELLS = 5_000_000

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_SS_NS = "urn:schemas-microsoft-com:office:spreadsheet"
# Excel 内置的日期 / 时间数字格式编号
_BUILTIN_DATE_FORMATS = set(range(14, 23)) | {27, 30, 36, 45, 46, 47, 50, 57}
_EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def is_excel(filename: str) -> bool:
    return filename.rsplit(".", 1)[-1].lower() in EXCEL_TYPES if "." in filename else False


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.replace(microsecond=0).isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    return str(value).strip()


def read_excel(file_bytes: bytes, sheet_name: str | None = None) -> tuple[list[list[str]], list[str], str]:
    """返回 (表格行, 全部工作表名, 实际读取的工作表名)。"""
    sheets = read_workbook(file_bytes)
    names = list(sheets)
    selected = sheet_name if sheet_name in sheets else names[0]
    rows = []
    for row in sheets[selected]:
        values = [_cell_text(v) for v in row]
        while values and values[-1] == "":
            values.pop()
        rows.append(values)
    if not any(any(c for c in r) for r in rows):
        raise ValueError(f"工作表「{selected}」没有数据")
    return rows, names, selected


def read_workbook(file_bytes: bytes) -> dict[str, list[list[Any]]]:
    """读取全部工作表的原始值（datetime / float / str / None），按格式自动选择解析方式。"""
    head = file_bytes[:2048]
    if zipfile.is_zipfile(io.BytesIO(file_bytes)):
        sheets = _read_xlsx(file_bytes)
    elif head.startswith(b"\xd0\xcf\x11\xe0"):
        sheets = _read_biff(file_bytes)
    elif b"<?xml" in head[:100] and _SS_NS.encode() in head:
        sheets = _read_spreadsheet_ml(file_bytes)
    else:
        sheets = _read_delimited(file_bytes)
    if not sheets:
        raise ValueError("Excel 文件没有工作表")
    return sheets


def _check_size(cells: int) -> None:
    if cells > MAX_CELLS:
        raise ValueError("Excel 数据量过大，请拆分后再导入")


# ---------------------------------------------------------------------------
# .xlsx
# ---------------------------------------------------------------------------


def _read_xlsx(file_bytes: bytes) -> dict[str, list[list[Any]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        # 部分系统导出的样式表不符合规范，openpyxl 拒绝打开；单元格数据本身是完好的
        return _read_xlsx_raw(file_bytes)
    try:
        result: dict[str, list[list[Any]]] = {}
        cells = 0
        for name in workbook.sheetnames[:MAX_SHEETS]:
            rows = []
            for row in workbook[name].iter_rows(values_only=True):
                rows.append(list(row))
                cells += len(row)
                _check_size(cells)
            result[name] = rows
        return result
    finally:
        workbook.close()


def _col_index(ref: str) -> int:
    index = 0
    for ch in ref:
        if not ch.isalpha():
            break
        index = index * 26 + (ord(ch.upper()) - 64)
    return index - 1


def _date_styles(archive: zipfile.ZipFile) -> set[int]:
    """返回使用日期格式的单元格样式下标（cellXfs 中的序号）。"""
    from defusedxml.ElementTree import fromstring

    try:
        root = fromstring(archive.read("xl/styles.xml"))
    except KeyError:
        return set()
    custom = {}
    fmts = root.find(f"{{{_MAIN_NS}}}numFmts")
    if fmts is not None:
        for fmt in fmts:
            code = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', "", fmt.get("formatCode", "")).lower()
            custom[int(fmt.get("numFmtId", "0"))] = bool(re.search(r"[ymdhs]", code))
    date_styles = set()
    xfs = root.find(f"{{{_MAIN_NS}}}cellXfs")
    for i, xf in enumerate(xfs if xfs is not None else []):
        fmt_id = int(xf.get("numFmtId", "0"))
        if fmt_id in _BUILTIN_DATE_FORMATS or custom.get(fmt_id):
            date_styles.add(i)
    return date_styles


def _from_serial(value: float) -> Any:
    if value < 1:  # 只有时间部分
        seconds = round(value * 86400)
        return dt.time(seconds // 3600 % 24, seconds // 60 % 60, seconds % 60)
    moment = _EXCEL_EPOCH + dt.timedelta(days=value)
    return moment.replace(microsecond=0) if moment.microsecond < 500_000 else (
        moment + dt.timedelta(seconds=1)).replace(microsecond=0)


def _read_xlsx_raw(file_bytes: bytes) -> dict[str, list[list[Any]]]:
    from defusedxml.ElementTree import fromstring

    try:
        archive = zipfile.ZipFile(io.BytesIO(file_bytes))
        workbook = fromstring(archive.read("xl/workbook.xml"))
        rels = fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    except Exception as exc:
        raise ValueError(f"Excel 文件无法解析: {exc}") from exc
    targets = {rel.get("Id"): rel.get("Target", "") for rel in rels.iter(f"{{{_PKG_REL_NS}}}Relationship")}
    shared: list[str] = []
    if "xl/sharedStrings.xml" in archive.namelist():
        for si in fromstring(archive.read("xl/sharedStrings.xml")).iter(f"{{{_MAIN_NS}}}si"):
            shared.append("".join(t.text or "" for t in si.iter(f"{{{_MAIN_NS}}}t")))
    date_styles = _date_styles(archive)

    result: dict[str, list[list[Any]]] = {}
    cells = 0
    sheets = workbook.find(f"{{{_MAIN_NS}}}sheets")
    for sheet in list(sheets if sheets is not None else [])[:MAX_SHEETS]:
        target = targets.get(sheet.get(f"{{{_REL_NS}}}id"), "").lstrip("/")
        path = target if target.startswith("xl/") else f"xl/{target}"
        try:
            root = fromstring(archive.read(path))
        except KeyError:
            continue
        rows: list[list[Any]] = []
        data = root.find(f"{{{_MAIN_NS}}}sheetData")
        for row in data if data is not None else []:
            row_no = int(row.get("r", len(rows) + 1))
            while len(rows) < row_no - 1:
                rows.append([])
            values: list[Any] = []
            for cell in row:
                ref = cell.get("r")
                col = _col_index(ref) if ref else len(values)
                while len(values) < col:
                    values.append(None)
                kind = cell.get("t", "n")
                raw = cell.find(f"{{{_MAIN_NS}}}v")
                text = raw.text if raw is not None else None
                if kind == "s" and text is not None:
                    value: Any = shared[int(text)]
                elif kind == "inlineStr":
                    value = "".join(t.text or "" for t in cell.iter(f"{{{_MAIN_NS}}}t"))
                elif kind == "b":
                    value = text == "1"
                elif kind in ("str", "e"):
                    value = text
                elif text is None:
                    value = None
                else:
                    number = float(text)
                    value = _from_serial(number) if int(cell.get("s", "0")) in date_styles else number
                values.append(value)
            rows.append(values)
            cells += len(values)
            _check_size(cells)
        result[sheet.get("name", f"Sheet{len(result) + 1}")] = rows
    return result


# ---------------------------------------------------------------------------
# XML 电子表格 2003（.xls 后缀）
# ---------------------------------------------------------------------------


def _read_spreadsheet_ml(file_bytes: bytes) -> dict[str, list[list[Any]]]:
    from defusedxml.ElementTree import fromstring

    try:
        root = fromstring(file_bytes)
    except Exception as exc:
        raise ValueError(f"XML 表格无法解析: {exc}") from exc
    ns = f"{{{_SS_NS}}}"
    result: dict[str, list[list[Any]]] = {}
    cells = 0
    for sheet in list(root.iter(f"{ns}Worksheet"))[:MAX_SHEETS]:
        rows: list[list[Any]] = []
        for row in sheet.iter(f"{ns}Row"):
            index = row.get(f"{ns}Index")
            while index and len(rows) < int(index) - 1:
                rows.append([])
            values: list[Any] = []
            for cell in row.iter(f"{ns}Cell"):
                col = cell.get(f"{ns}Index")
                while col and len(values) < int(col) - 1:
                    values.append(None)
                data = cell.find(f"{ns}Data")
                value: Any = None
                if data is not None:
                    text = "".join(data.itertext()).strip()
                    kind = data.get(f"{ns}Type", "String")
                    if kind == "Number" and text:
                        value = float(text)
                    elif kind == "DateTime" and text:
                        value = dt.datetime.fromisoformat(text[:19])
                    else:
                        value = text
                values.append(value)
            rows.append(values)
            cells += len(values)
            _check_size(cells)
        result[sheet.get(f"{ns}Name", f"Sheet{len(result) + 1}")] = rows
    return result


# ---------------------------------------------------------------------------
# 制表符 / 逗号分隔的文本（.xls 后缀）与二进制 .xls
# ---------------------------------------------------------------------------


def _read_delimited(file_bytes: bytes) -> dict[str, list[list[Any]]]:
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    first = text.split("\n", 1)[0] if text else ""
    if text is None or "\x00" in text[:2048] or ("\t" not in first and "," not in first):
        raise ValueError("文件不是可识别的 Excel 格式（请在 Excel 中另存为 .xlsx 后再导入）")
    delimiter = "\t" if first.count("\t") >= first.count(",") else ","
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        while row and not row[-1].strip():
            row.pop()
        rows.append([cell.strip() or None for cell in row])
    _check_size(sum(len(r) for r in rows))
    if not any(rows):
        raise ValueError("文件没有数据")
    return {"Sheet1": rows}


def _read_biff(file_bytes: bytes) -> dict[str, list[list[Any]]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - xlrd 在 requirements 中
        raise ValueError("旧版 .xls 请先在 Excel 中另存为 .xlsx") from exc
    try:
        book = xlrd.open_workbook(file_contents=file_bytes)
    except Exception as exc:
        raise ValueError(f"旧版 .xls 文件无法解析: {exc}") from exc
    result: dict[str, list[list[Any]]] = {}
    cells = 0
    for sheet in book.sheets()[:MAX_SHEETS]:
        rows = []
        for r in range(sheet.nrows):
            values: list[Any] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    values.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    values.append(None)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    values.append(bool(cell.value))
                else:
                    values.append(cell.value)
            rows.append(values)
            cells += len(values)
            _check_size(cells)
        result[sheet.name] = rows
    return result
