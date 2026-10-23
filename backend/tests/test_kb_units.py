import io

import pytest

from backend.auth import hash_password, validate_password_strength, verify_password
from backend.kb.chunker import split_markdown
from backend.kb.converter import convert_to_markdown, file_type_of
from backend.kb.retrieval import make_snippet, query_terms
from backend.kb.wiki import extract_links, slugify, summarize


def _docx_bytes():
    from docx import Document

    doc = Document()
    doc.add_heading("报销制度", 0)
    doc.add_heading("住宿", 1)
    doc.add_paragraph("一线城市每晚 600 元。")
    doc.add_paragraph("提前申请", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "级别", "餐补"
    table.cell(1, 0).text, table.cell(1, 1).text = "经理", "150|元"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_keeps_structure():
    doc = convert_to_markdown("报销.docx", _docx_bytes())
    assert doc.title == "报销制度"
    assert "## 住宿" in doc.markdown and "- 提前申请" in doc.markdown
    assert "| 级别 | 餐补 |" in doc.markdown and "150\\|元" in doc.markdown


def test_txt_escapes_markdown_and_gbk():
    doc = convert_to_markdown("说明.txt", "# 不是标题\n正文".encode("gbk"))
    assert doc.title == "说明"
    assert "\\# 不是标题" in doc.markdown


def test_unsupported_and_empty_files():
    with pytest.raises(ValueError):
        file_type_of("a.xlsx")
    with pytest.raises(ValueError):
        convert_to_markdown("a.md", b"   ")


def test_chunker_tracks_heading_path_and_size():
    md = "# 手册\n\n## 第一章\n\n" + ("这是一句话。" * 200) + "\n\n## 第二章\n\n短段落\n\n```sql\nSELECT 1;\n```\n"
    chunks = split_markdown(md, chunk_size=300, overlap=50)
    assert all(len(c.content) <= 360 for c in chunks)
    assert chunks[0].heading_path == "手册 > 第一章"
    assert chunks[-1].heading_path == "手册 > 第二章" and "SELECT 1;" in chunks[-1].content


def test_query_terms_and_snippet():
    terms = query_terms("一线城市住宿标准是什么？ GMV")
    assert "住宿" in terms and "GMV" in terms and "什么" not in terms
    assert "住宿" in make_snippet("前文" * 100 + "住宿标准 600 元", ["住宿"])


def test_wiki_helpers():
    assert slugify("客户服务 SLA") == "客户服务-sla"
    assert extract_links("见 [[客户服务 SLA]] 和 [[metrics|指标]]") == {"客户服务-sla", "metrics"}
    assert summarize("# 标题\n\n**重点** [[metrics|指标口径]] 内容") == "重点 指标口径 内容"


def test_password_hashing():
    stored = hash_password("Secret123")
    assert verify_password("Secret123", stored) and not verify_password("secret123", stored)
    with pytest.raises(ValueError):
        validate_password_strength("12345678")
